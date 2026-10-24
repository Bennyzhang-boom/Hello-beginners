import json
import os
import time
from collections import defaultdict, deque
from pathlib import Path
from typing import Any

import fitz
import openpyxl
from docx import Document
from fastapi import FastAPI, File, HTTPException, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from openai import OpenAI
from pptx import Presentation

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_FILES = 8
MAX_TOTAL_CHARS = 120_000
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".xlsx", ".txt", ".md"}
MODEL = os.getenv("OPENAI_MODEL", "gpt-4.1-mini")
API_KEY = os.getenv("OPENAI_API_KEY")

app = FastAPI(title="Huizhi Document API", version="1.0.0")
origins = [x.strip() for x in os.getenv(
    "FRONTEND_ORIGINS",
    "https://bennyzhang-boom.github.io,http://localhost:8000,http://127.0.0.1:8000"
).split(",") if x.strip()]
app.add_middleware(
    CORSMiddleware,
    allow_origins=origins,
    allow_credentials=False,
    allow_methods=["GET", "POST", "OPTIONS"],
    allow_headers=["Content-Type"],
)

# Basic in-memory guardrail for the prototype. Use a persistent limiter/authentication
# before opening the service to a broad public audience.
_request_times: dict[str, deque[float]] = defaultdict(deque)


def check_rate_limit(request: Request) -> None:
    now = time.time()
    bucket = _request_times[request.client.host if request.client else "unknown"]
    while bucket and now - bucket[0] > 3600:
        bucket.popleft()
    if len(bucket) >= 5:
        raise HTTPException(status_code=429, detail="请求过于频繁，请一小时后再试。")
    bucket.append(now)


def extract_pdf(data: bytes) -> tuple[str, int]:
    chunks = []
    with fitz.open(stream=data, filetype="pdf") as pdf:
        for page_num, page in enumerate(pdf, start=1):
            text = page.get_text("text").strip()
            if text:
                chunks.append(f"[第 {page_num} 页]\n{text}")
    return "\n\n".join(chunks), len(chunks)


def extract_docx(data: bytes) -> tuple[str, int]:
    import io
    document = Document(io.BytesIO(data))
    chunks = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table_index, table in enumerate(document.tables, start=1):
        rows = []
        for row in table.rows:
            values = [cell.text.strip().replace("\n", " / ") for cell in row.cells]
            if any(values):
                rows.append(" | ".join(values))
        if rows:
            chunks.append(f"[表格 {table_index}]\n" + "\n".join(rows))
    return "\n".join(chunks), len(chunks)


def extract_pptx(data: bytes) -> tuple[str, int]:
    import io
    deck = Presentation(io.BytesIO(data))
    chunks = []
    for slide_num, slide in enumerate(deck.slides, start=1):
        items = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                items.append(shape.text.strip())
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    items.append(" | ".join(cell.text.strip() for cell in row.cells))
        if items:
            chunks.append(f"[第 {slide_num} 张幻灯片]\n" + "\n".join(items))
    return "\n\n".join(chunks), len(chunks)


def extract_xlsx(data: bytes) -> tuple[str, int]:
    import io
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    chunks = []
    row_count = 0
    try:
        for sheet in workbook.worksheets:
            lines = []
            for row_num, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                values = ["" if value is None else str(value) for value in row]
                if any(value.strip() for value in values):
                    lines.append(f"第{row_num}行: " + " | ".join(values))
                    row_count += 1
                    if row_count > 10000:
                        raise HTTPException(status_code=413, detail="Excel 数据行数过多，请拆分后再上传。")
            if lines:
                chunks.append(f"[工作表：{sheet.title}]\n" + "\n".join(lines))
    finally:
        workbook.close()
    return "\n\n".join(chunks), row_count


def extract_text(data: bytes, ext: str) -> tuple[str, int]:
    if ext == ".pdf":
        return extract_pdf(data)
    if ext == ".docx":
        return extract_docx(data)
    if ext == ".pptx":
        return extract_pptx(data)
    if ext == ".xlsx":
        return extract_xlsx(data)
    try:
        return data.decode("utf-8-sig"), 1
    except UnicodeDecodeError:
        raise HTTPException(status_code=400, detail="TXT/MD 文件必须使用 UTF-8 编码。")


@app.get("/")
def root() -> dict[str, str]:
    return {"service": "Huizhi Document API", "status": "ok", "docs": "/docs"}


@app.get("/health")
def health() -> dict[str, str]:
    return {"status": "ok"}


@app.post("/api/analyze")
async def analyze(request: Request, files: list[UploadFile] = File(...)) -> dict[str, Any]:
    check_rate_limit(request)
    if not API_KEY:
        raise HTTPException(status_code=503, detail="后端尚未配置 OPENAI_API_KEY 环境变量。")
    if not files or len(files) > MAX_FILES:
        raise HTTPException(status_code=400, detail=f"请一次上传 1 至 {MAX_FILES} 个文件。")

    parsed_files = []
    total_chars = 0
    for upload in files:
        name = Path(upload.filename or "unnamed").name
        ext = Path(name).suffix.lower()
        if ext not in ALLOWED_EXTENSIONS:
            raise HTTPException(status_code=415, detail=f"暂不支持 {ext or '未知格式'}：{name}。请使用 PDF、DOCX、PPTX、XLSX、TXT 或 MD。")
        data = await upload.read(MAX_FILE_BYTES + 1)
        if not data:
            raise HTTPException(status_code=400, detail=f"文件为空：{name}")
        if len(data) > MAX_FILE_BYTES:
            raise HTTPException(status_code=413, detail=f"单个文件不能超过 20 MB：{name}")
        try:
            extracted, units = extract_text(data, ext)
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(status_code=400, detail=f"无法解析文件 {name}。文件可能已损坏或格式不匹配。") from exc
        if not extracted.strip():
            hint = "（若这是扫描版 PDF，需要先做 OCR）" if ext == ".pdf" else ""
            raise HTTPException(status_code=422, detail=f"未能从 {name} 提取到可用文字{hint}。")
        total_chars += len(extracted)
        if total_chars > MAX_TOTAL_CHARS:
            raise HTTPException(status_code=413, detail="所有文件提取的文字合计超过 120,000 字符，请分批整理。")
        parsed_files.append({
            "name": name,
            "text": extracted,
            "units": units,
            "extracted_chars": len(extracted),
        })

    document_text = "\n\n".join(
        f"===== 文件：{item['name']} =====\n{item['text']}" for item in parsed_files
    )
    instructions = (
        "你是中文文档整理助手。仅根据用户提供的文档内容进行归纳，不要编造事实。"
        "请用简体中文返回 JSON 对象，字段必须为：title（简短标题字符串）、"
        "overview（约 150-300 字概述）、key_points（3-8 条关键结论的字符串数组）、"
        "keywords（5-12 个关键词的字符串数组）、sources（数组，每项含 file、locator、evidence，"
        "locator 尽量引用文档内的页码、幻灯片、工作表或行号，evidence 为简短原文依据）。"
        "若文档没有足够证据，明确说明，不要虚构引用。"
    )
    prompt = (
        "请整理以下文件内容。文件名和方括号中的页码、幻灯片、工作表、行号是来源定位信息。"
        "把跨文件共识、差异和待确认事项说清楚。\n\n" + document_text
    )
    try:
        client = OpenAI(api_key=API_KEY)
        response = client.responses.create(
            model=MODEL,
            instructions=instructions,
            input=prompt,
            text={"format": {"type": "json_object"}},
        )
        payload = json.loads(response.output_text)
    except json.JSONDecodeError as exc:
        raise HTTPException(status_code=502, detail="AI 返回格式无法解析，请稍后重试。") from exc
    except Exception as exc:
        raise HTTPException(status_code=502, detail="调用 OpenAI API 失败，请检查后端日志、模型配置和 API 额度。") from exc

    return {
        "title": str(payload.get("title") or "文档整理结果")[:120],
        "overview": str(payload.get("overview") or ""),
        "key_points": [str(x) for x in payload.get("key_points", []) if str(x).strip()][:12],
        "keywords": [str(x) for x in payload.get("keywords", []) if str(x).strip()][:20],
        "sources": payload.get("sources", []) if isinstance(payload.get("sources", []), list) else [],
        "files": [{"name": x["name"], "extracted_chars": x["extracted_chars"], "units": x["units"]} for x in parsed_files],
        "model": MODEL,
    }
